"""
Report writer.

Instead of rebuilding an Excel file from scratch, this makes a literal
copy of the golden dataset file and appends result columns onto that
copy. Two reasons this matters:

  1. The master golden file is never opened in write mode -- it's only
     ever read (by data_loader.py) or copied (here). No execution can
     ever corrupt/overwrite your source of truth.
  2. Any extra columns you keep in the golden sheet for your own
     bookkeeping (notes, owner, priority, whatever) survive into the
     report untouched, instead of being dropped because EvalRow doesn't
     know about them.

The copy becomes this run's report: golden columns stay as-is, judge/
scoring columns are appended after them.
"""
import shutil
from openpyxl import load_workbook
from src.models import EvalRow

# Columns appended to the golden sheet's own columns, in this order.
# 'actual' and 'report' come first so the sheet reads: id, query, expected,
# module, actual, report, ... -- the columns someone scanning results cares
# about most, right next to the golden data they're being compared against.
RESULT_COLUMNS = [
    "actual", "report",
    "judge_score", "judge_pass", "judge_reason",
    "wer", "latency_ms", "final_pass",
    "run_id", "timestamp", "error",
]


def write_report(rows: list[EvalRow], golden_path: str, output_path: str) -> float:
    """
    Copies golden_path -> output_path, then edits the copy in place:
    appends RESULT_COLUMNS as new columns and fills in one row per
    test case, matched by the 'id' column. Returns the pass rate (%).
    """
    # 1. Copy the golden file -- the master is never touched.
    shutil.copy(golden_path, output_path)

    # 2. Open the copy for editing.
    wb = load_workbook(output_path)
    ws = wb.active

    # Map existing header names -> column index (1-based, openpyxl convention).
    # Matched case-insensitively -- real golden sheets aren't all lowercase
    # (a hand-built demo sheet uses "id", a real question bank export might
    # use "ID" instead), same as data_loader.py's column matching.
    header_row = 1
    existing_headers = {
        cell.value: cell.column for cell in ws[header_row] if cell.value is not None
    }
    existing_headers_lower = {
        str(name).strip().lower(): col for name, col in existing_headers.items()
    }
    if "id" not in existing_headers_lower:
        raise ValueError(f"Golden sheet at {golden_path} has no 'id' column to match results against")
    id_col = existing_headers_lower["id"]

    # 3. Append new headers after the last existing column.
    next_col = ws.max_column + 1
    result_col_index = {}
    for offset, col_name in enumerate(RESULT_COLUMNS):
        col_idx = next_col + offset
        ws.cell(row=header_row, column=col_idx, value=col_name)
        result_col_index[col_name] = col_idx

    # 4. Index results by test case id for fast lookup while walking rows.
    results_by_id = {r.id: r for r in rows}

    passed_count = 0
    total_matched = 0
    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_id = str(ws.cell(row=row_idx, column=id_col).value)
        result = results_by_id.get(row_id)
        if result is None:
            continue  # a golden row that wasn't run this time -- leave result cells blank

        total_matched += 1
        if result.final_pass:
            passed_count += 1

        values = result.model_dump()
        for col_name in RESULT_COLUMNS:
            ws.cell(row=row_idx, column=result_col_index[col_name], value=values.get(col_name))

    wb.save(output_path)

    pass_rate = (passed_count / total_matched * 100) if total_matched else 0.0
    return pass_rate
